"""
report_docx.py
==============
Builds the KeRRA "Daily Progress Report / Site Diary" as a Word (.docx) file,
laid out like the sample report (logo + title on every page, sections I to X, page number footer).
"""

import io
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from PIL import Image, ImageOps

import report_data as rd

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Times New Roman"
GREEN = RGBColor(0x00, 0xB0, 0x50)
CONTENT_W = 6.27                                  # inches (A4 less 1" margins)
C, L, R = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT


# ------------------------------------------------------------------ low-level helpers
def _font(run, size=12, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _fmt(p, align=None, before=0, after=0, keep_next=False):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if align is not None:
        p.alignment = align
    if keep_next:
        pf.keep_with_next = True


def _field(par, code, size=10, name="Arial", cached="1", italic=None, color=None):
    """A Word field (PAGE, SEQ ...). Each part is its own run so the font applies to the displayed result."""
    def run_with(child):
        r = par.add_run()
        _font(r, size, name=name, italic=italic, color=color)
        r._r.append(child)
    for kind in ("begin", "code", "separate", "text", "end"):
        if kind in ("begin", "separate", "end"):
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        elif kind == "code":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = f" {code} "
        else:
            e = OxmlElement("w:t")
            e.text = cached
        run_with(e)


def _border_bottom(p, val="thinThickSmallGap", sz=24):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    for k, v in (("val", val), ("sz", str(sz)), ("space", "1"), ("color", "auto")):
        e.set(qn(f"w:{k}"), v)
    b.append(e)
    pPr.insert_element_before(b, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr")


def _set_cell(cell, text="", size=11, bold=False, italic=False, align=L, width=None):
    cell.text = ""
    lines = str(text).split("\n") if text != "" else [""]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        _fmt(p, align, 1, 1)
        if line:
            _font(p.add_run(line), size, bold, italic)
    if width is not None:
        cell.width = Inches(width)
    return cell


def _table(doc, widths, rows=0):
    t = doc.add_table(rows=rows, cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    return t


def _add_row(t, widths, header=False):
    row = t.add_row()
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)
    if header:
        h = OxmlElement("w:tblHeader")
        trPr.append(h)
    for c, w in zip(row.cells, widths):
        c.width = Inches(w)
    return row


def _heading(doc, numeral, text):
    p = doc.add_paragraph()
    _fmt(p, L, 12, 4, keep_next=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5), WD_TAB_ALIGNMENT.LEFT)
    _font(p.add_run(f"{numeral}.\t{text}"), 12, bold=True)
    return p


def _gap(doc, pts=4):
    p = doc.add_paragraph()
    _fmt(p, L, 0, 0)
    p.paragraph_format.line_spacing = Pt(pts)
    return p


def _grid_table(doc, widths, headers, rows, size=10.5, aligns=None, min_rows=0, number=True):
    """Generic table: header row + numbered data rows, padded with empty rows up to min_rows."""
    aligns = aligns or [L] * len(headers)
    t = _table(doc, widths)
    hr = _add_row(t, widths, header=True)
    for c, h, w, a in zip(hr.cells, headers, widths, aligns):
        _set_cell(c, h, size, True, align=C if a == C else L, width=w)
    n = max(len(rows), min_rows)
    for i in range(n):
        vals = rows[i] if i < len(rows) else [""] * (len(headers) - 1)
        r = _add_row(t, widths)
        cells = ([f"{i + 1}."] if number else []) + [str(v) for v in vals]
        for c, v, w, a in zip(r.cells, cells, widths, aligns):
            _set_cell(c, v, size, align=a, width=w)
    return t


# ------------------------------------------------------------------ photo handling
PHOTO_W, PHOTO_H = 2.9, 2.175          # inches - every photo is cut to this same 4:3 size
CAPTION_W = 2.2                         # inches - width of the caption block
CAPTION_BLUE = RGBColor(0x1F, 0x49, 0x7D)
PHOTO_ROW_MIN_H = 3.1                   # inches - minimum row height (as in the sample report)


def _prepare_photo(raw, px_w=1200):
    """Turn any photo into the same-size picture: upright (EXIF), centre-cropped to 4:3, 1200 x 900 px."""
    img = ImageOps.exif_transpose(Image.open(io.BytesIO(raw))).convert("RGB")
    img = ImageOps.fit(img, (px_w, round(px_w * PHOTO_H / PHOTO_W)), method=Image.LANCZOS, centering=(0.5, 0.5))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    buf.seek(0)
    return buf


def _photo_cell(cell, raw, caption, number):
    cell.text = ""
    p = cell.paragraphs[0]
    _fmt(p, C, 2, 4, keep_next=True)
    p.add_run().add_picture(_prepare_photo(raw), width=Inches(PHOTO_W), height=Inches(PHOTO_H))
    # caption: the block is centred in the cell, the text inside it is left aligned
    cp = cell.add_paragraph()
    _fmt(cp, L, 0, 4)
    side = (PHOTO_W - CAPTION_W) / 2 + 0.04
    cp.paragraph_format.left_indent = Inches(side)
    cp.paragraph_format.right_indent = Inches(side)
    _font(cp.add_run("Figure "), 9, italic=True, color=CAPTION_BLUE, name="Calibri")
    _field(cp, "SEQ Figure \\* ARABIC", size=9, name="Calibri", cached=str(number), italic=True, color=CAPTION_BLUE)
    _font(cp.add_run(f": {caption}"), 9, italic=True, color=CAPTION_BLUE, name="Calibri")


def _photo_table(doc, photos):
    w = [CONTENT_W / 2, CONTENT_W / 2]
    t = _table(doc, w)
    n = 0
    for start in range(0, len(photos), 2):                 # a new row only when there are more photos
        row = _add_row(t, w)
        trPr = row._tr.get_or_add_trPr()
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), str(int(PHOTO_ROW_MIN_H * 1440)))
        h.set(qn("w:hRule"), "atLeast")
        trPr.append(h)
        for ci in range(2):
            cell = row.cells[ci]
            if start + ci < len(photos):
                n += 1
                raw, cap = photos[start + ci]
                _photo_cell(cell, raw, cap, n)
            else:
                cell.text = ""
            cell.width = Inches(w[ci])


# ------------------------------------------------------------------ the report
def build_report(st, photos=None, logo_path=None):
    """st: dict from report_data (as held by the form).  photos: list of (bytes, caption).  Returns .docx bytes."""
    logo_path = logo_path or os.path.join(HERE, "logo.png")
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Emu(7560310), Emu(10692130)          # A4
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.49)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for z in doc.settings.element.findall(qn("w:zoom")):
        z.set(qn("w:percent"), "100")
    doc.core_properties.title = "Daily Progress Report / Site Diary"
    doc.core_properties.author = st.get("contractor", "") or "KeRRA"

    # ---- header (every page): logo, authority name, report title with double rule
    hdr = sec.header
    hp = hdr.paragraphs[0]
    _fmt(hp, C, 0, 0)
    if os.path.exists(logo_path):
        hp.add_run().add_picture(logo_path, width=Inches(1.42))
    p = hdr.add_paragraph()
    _fmt(p, C, 6, 6)
    _font(p.add_run("KENYA RURAL ROADS AUTHORITY"), 22, bold=True, color=GREEN)
    p = hdr.add_paragraph()
    _fmt(p, C, 0, 10)
    _font(p.add_run("DAILY PROGRESS REPORT/SITE DIARY"), 18, bold=True)
    _border_bottom(p)

    # ---- footer: Page | n
    fp = sec.footer.paragraphs[0]
    _fmt(fp, L)
    _font(fp.add_run("Page | "), 10, name="Arial")
    _field(fp, "PAGE")


    # ---- project block
    def kv(label, value, gap=6):
        p = doc.add_paragraph()
        _fmt(p, L, 0, gap)
        _font(p.add_run(label + " "), 12, bold=True)
        _font(p.add_run(value or ""), 12)

    kv("PROJECT NAME:", (st["project"] or "").upper())
    kv("CONTRACT NO.:", st["contract_no"])
    kv("CONTRACTOR:", (st["contractor"] or "").upper())
    p = doc.add_paragraph()
    _fmt(p, L, 0, 6)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.2), WD_TAB_ALIGNMENT.LEFT)
    _font(p.add_run("Date:\t"), 12, bold=True)
    _font(p.add_run(st["date"].strftime("%d/%m/%Y")), 12, bold=True)

    # ---- I. weather
    _heading(doc, "I", "WEATHER CONDITIONS")
    w = [1.75, 0.95, 3.57]
    t = _table(doc, w)
    station = f"{st['station_from'] or ''} TO {st['station_to'] or ''}".strip() if (st["station_from"] or st["station_to"]) else ""
    rain = [f"{r[0]}:{r[1]} mm" for r in rd.clean_rows(st["rainfall"]) if r[1] != ""]
    rows = [("STATION", station), ("MORNING", st["morning"]), ("AFTERNOON", st["afternoon"]),
            ("RAINFALL (mm)", "\n".join(rain) if rain else "Nil"), ("GROUND CONDITIONS", st["ground"])]
    for label, val in rows:
        r = _add_row(t, w)
        a = r.cells[0].merge(r.cells[1])
        _set_cell(a, label, 11, True)
        _set_cell(r.cells[2], val, 11)
    for sub, key in (("MAX (°C)", "tmax"), ("MIN (°C)", "tmin")):
        r = _add_row(t, w)
        _set_cell(r.cells[1], sub, 11, True, italic=False)
        _set_cell(r.cells[2], rd.fmt_num(st[key]), 11)
    top = t.rows[-2].cells[0].merge(t.rows[-1].cells[0])
    _set_cell(top, "TEMP.", 11, True)

    # ---- II. activities
    _heading(doc, "II", "ACTIVITIES")
    w = [0.45, 1.3, 0.8, 0.8, 0.6, 0.7, 0.65, 0.97]
    t = _table(doc, w)
    h1, h2 = _add_row(t, w, True), _add_row(t, w, True)
    for i in (0, 1, 7):
        m = h1.cells[i].merge(h2.cells[i])
        _set_cell(m, ["S/NO.", "OPERATION", "", "", "", "", "", "REMARK / STATUS"][i], 10, True, width=w[i])
    _set_cell(h1.cells[2].merge(h1.cells[3]), "CHAINAGE", 10, True, align=C)
    _set_cell(h1.cells[4].merge(h1.cells[6]), "LENGTH (m)", 10, True, align=C)
    for i, txt in zip(range(2, 7), ["FROM", "TO", "PREV.", "TODAY", "TOTAL"]):
        _set_cell(h2.cells[i], txt, 9.5, True, align=C, width=w[i])
    acts = rd.activity_rows(st["activities"])
    for i in range(max(len(acts), 2)):
        r = _add_row(t, w)
        a = acts[i] if i < len(acts) else None
        vals = [f"{i + 1}" if a else "", a["op"] if a else "", a["from_txt"] if a else "", a["to_txt"] if a else "",
                rd.fmt_num(a["prev"]) if a else "", rd.fmt_num(a["today"]) if a else "",
                rd.fmt_num(a["total"]) if a else "", a["remark"] if a else ""]
        for c, v, ww, al in zip(r.cells, vals, w, [R, L, C, C, C, C, C, L]):
            _set_cell(c, v, 10.5, align=al, width=ww)

    # ---- III. equipment
    _heading(doc, "III", "CONTRACTOR’S EQUIPMENT /MACHINERY ON SITE")
    w = [0.45, 1.35, 0.45, 1.0, 0.7, 0.7, 0.65, 0.97]
    t = _table(doc, w)
    h1, h2 = _add_row(t, w, True), _add_row(t, w, True)
    for i, txt in ((0, "S/NO."), (1, "EQUIPMENT / MACHINERY"), (2, "QTY"), (3, "REG. NO."), (6, "IDLE HOURS"), (7, "STATUS")):
        _set_cell(h1.cells[i].merge(h2.cells[i]), txt, 9.5, True, align=C if i in (0, 2, 6) else L, width=w[i])
    _set_cell(h1.cells[4].merge(h1.cells[5]), "HOURS DEPLOYED", 9.5, True, align=C)
    _set_cell(h2.cells[4], "From:", 9.5, True, align=C, width=w[4])
    _set_cell(h2.cells[5], "To:", 9.5, True, align=C, width=w[5])
    eq = rd.clean_rows(st["equipment"])
    for i, row in enumerate(eq):
        r = _add_row(t, w)
        for c, v, ww, al in zip(r.cells, [f"{i + 1}."] + row, w, [R, L, C, L, C, C, C, L]):
            _set_cell(c, v, 10, align=al, width=ww)
    if not eq:
        r = _add_row(t, w)
        for c, ww in zip(r.cells, w):
            _set_cell(c, "", 10, width=ww)

    # ---- IV. visitors
    _heading(doc, "IV", "VISITORS")
    lines = [x for x in (st["visitors"] or "Nil").split("\n") if x.strip()] or ["Nil"]
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        _fmt(p, L, 0, 0)
        _font(p.add_run(line.strip()), 12, italic=len(lines) == 1 and line.strip().lower() in ("nil", "ni", "none"))

    # ---- V. labour
    _heading(doc, "V", "CONTRACTOR’S LABOUR FORCE")
    w = [0.5, 1.9, 0.9, 0.9, 0.9, 1.17]
    t = _table(doc, w)
    h1, h2 = _add_row(t, w, True), _add_row(t, w, True)
    for i, txt in ((0, "S/NO."), (1, "CATEGORY"), (2, "NUMBER"), (5, "IDLE HOURS")):
        _set_cell(h1.cells[i].merge(h2.cells[i]), txt, 11, True, align=C if i in (0, 2, 5) else L, width=w[i])
    _set_cell(h1.cells[3].merge(h1.cells[4]), "TIME DEPLOYED", 11, True, align=C)
    _set_cell(h2.cells[3], "From:", 11, True, align=C, width=w[3])
    _set_cell(h2.cells[4], "To:", 11, True, align=C, width=w[4])
    lab = rd.clean_rows(st["labour"])
    for i, row in enumerate(lab):
        r = _add_row(t, w)
        for c, v, ww, al in zip(r.cells, [f"{i + 1}."] + row, w, [R, L, C, C, C, C]):
            _set_cell(c, v, 11, align=al, width=ww)
    if lab:
        r = _add_row(t, w)
        _set_cell(r.cells[0].merge(r.cells[1]), "TOTAL", 11, True, align=R)
        _set_cell(r.cells[2], rd.fmt_num(sum(rd.to_float(x[1]) or 0 for x in lab)), 11, True, align=C)
        for c in r.cells[3:]:
            _set_cell(c, "", 11)

    # ---- VI. materials
    _heading(doc, "VI", "MATERIALS DELIVERED")
    mats = rd.clean_rows(st["materials"])
    w = [3.13, 3.14]
    t = _table(doc, w)
    r = _add_row(t, w, True)
    _set_cell(r.cells[0], "MATERIAL", 11, True, align=C, width=w[0])
    _set_cell(r.cells[1], "TOTAL QUANTITY DELIVERED", 11, True, align=C, width=w[1])
    for i in range(max(len(mats), 3)):
        r = _add_row(t, w)
        v = mats[i] if i < len(mats) else ["", "", ""]
        _set_cell(r.cells[0], v[0], 11, width=w[0])
        _set_cell(r.cells[1], f"{v[1]} {v[2]}".strip(), 11, width=w[1])

    # ---- VII. accidents, delays & problems
    _heading(doc, "VII", "ACCIDENTS, DELAYS & PROBLEMS")
    for line in [x.strip() for x in (st["issues"] or rd.NO_DELAYS).split("\n") if x.strip()] or [rd.NO_DELAYS]:
        p = doc.add_paragraph()
        _fmt(p, L, 0, 2)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5), WD_TAB_ALIGNMENT.LEFT)
        _font(p.add_run("➢"), 11, name="Segoe UI Symbol")
        _font(p.add_run("\t" + line), 12)

    # ---- VIII. RE vehicles
    _heading(doc, "VIII", "RESIDENT ENGINEER’S SUPERVISION VEHICLES")
    _grid_table(doc, [0.55, 1.1, 1.0, 1.2, 1.2, 1.22],
                ["S/NO.", "VEHICLE REG. NO.", "ASSIGNED", "INSURANCE STATUS", "DAILY CONDITION", "REMARKS"],
                rd.clean_rows(st["vehicles"]), 10.5, [R, L, L, L, C, L], min_rows=3)

    # ---- IX. RE staff activities
    _heading(doc, "IX", "RESIDENT ENGINEER’S STAFF ACTIVITIES")
    _grid_table(doc, [0.55, 1.5, 4.22], ["S/NO.", "TEAM", "ACTIVITY"],
                rd.clean_rows(st["staff"]), 11, [R, L, L], min_rows=3)

    # ---- X. photos: two per row, centred in bordered cells, "Figure n: caption" under each photo
    photos = photos or []
    _heading(doc, "X", "PROGRESS PHOTO" + ("S" if len(photos) > 1 else ""))
    if photos:
        _photo_table(doc, photos)
    else:
        p = doc.add_paragraph()
        _fmt(p, L)
        _font(p.add_run("No photos attached."), 12, italic=True)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def report_filename(st):
    return f"Daily report - {st['date'].strftime('%d-%m-%Y')}.docx"
